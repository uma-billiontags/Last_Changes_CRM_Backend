import io
import json
import re
from datetime import timedelta, datetime

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from django.core.files.base import ContentFile
from django.http import FileResponse, HttpResponse

from rest_framework.decorators import api_view, parser_classes
from rest_framework.parsers import MultiPartParser
from rest_framework.response import Response

from django.shortcuts import get_object_or_404
from django.http import JsonResponse
from django.views.decorators.csrf import csrf_exempt
from django.core.mail import EmailMultiAlternatives
from django.conf import settings

from campaigns.models import Campaign, LineItem
from .models import CampaignDailyEntry, CampaignDailyReportExcel


# ── Helpers ──────────────────────────────────────────────────────────────────

def parse_date(date_str):
    if not date_str:
        return None
    return datetime.fromisoformat(date_str.replace("Z", "")).date()


def _parse_geo(geo_targeting):
    """Parse geo_targeting field (JSON string or list) into readable string."""
    try:
        if isinstance(geo_targeting, str):
            data = json.loads(geo_targeting)
        else:
            data = geo_targeting
        parts = []
        for loc in (data if isinstance(data, list) else [data]):
            segment = ", ".join(
                filter(None, [loc.get("country", ""), loc.get("state", ""), loc.get("city", "")])
            )
            if segment:
                parts.append(segment)
        return " | ".join(parts) if parts else ""
    except Exception:
        return str(geo_targeting) if geo_targeting else ""


def daterange(start_date, end_date):
    days = (end_date - start_date).days
    for i in range(days + 1):
        yield start_date + timedelta(days=i)


def _safe_sheet_name(name: str, used_names: set) -> str:
    """Excel sheet names: max 31 chars, no : \\ / ? * [ ]. Dedupe within a workbook."""
    cleaned = re.sub(r"[:\\/?*\[\]]", "-", str(name)).strip() or "Sheet"
    base = cleaned[:31]
    candidate = base
    suffix = 1
    while candidate in used_names:
        suffix_str = f"~{suffix}"
        candidate = base[: 31 - len(suffix_str)] + suffix_str
        suffix += 1
    used_names.add(candidate)
    return candidate


def _is_video_line_item(li) -> bool:
    """Detect video line items from the raw ad_format JSONField (list or string)."""
    raw = li.ad_format
    if isinstance(raw, list):
        text = " ".join(str(x) for x in raw)
    else:
        text = str(raw or "")
    return "video" in text.lower() or "youtube" in text.lower()


# ── Build the Daily Reports Excel workbook ──────────────────────────────────

def build_daily_report_excel(campaign) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    header_font = Font(name="Arial", bold=True, size=10)
    value_font = Font(name="Arial", size=10)
    label_fill = PatternFill("solid", start_color="D9E1F2")
    border_side = Side(style="thin", color="B8CCE4")
    thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
    table_header_fill = PatternFill("solid", start_color="1F4E79")
    table_header_font = Font(name="Arial", bold=True, size=10, color="FFFFFF")

    insertion_order = getattr(campaign, "insertion_order", None)
    # ✅ FIX: this schema has no campaign_id field — fall back to ticket_id
    main_sheet_name = (
        str(insertion_order.io_id)[:31] if insertion_order else str(campaign.ticket_id)[:31]
    )

    used_sheet_names = {main_sheet_name}
    ws_main = wb.create_sheet(title=main_sheet_name)

    line_items = list(campaign.line_items.all())

    entries = CampaignDailyEntry.objects.filter(campaign=campaign).select_related("line_item")
    entries_map = {}
    for e in entries:
        entries_map.setdefault(e.line_item_id, {})[e.entry_date] = e

    video_line_items = [li for li in line_items if _is_video_line_item(li)]
    non_video_line_items = [li for li in line_items if not _is_video_line_item(li)]

    def write_summary_header(ws, start_col, li, target_impressions, first_data_row,
                              last_data_row, col_letters, daily_target_col_idx):
        col_d = col_letters[daily_target_col_idx]

        ws.cell(row=3, column=start_col, value=li.start_date.strftime("%d %B, %Y"))
        ws.merge_cells(start_row=3, start_column=start_col + 1, end_row=3, end_column=start_col + 2)
        ws.cell(row=3, column=start_col + 1, value=li.line_item_name or li.line_item_id)
        ws.cell(row=3, column=start_col + 1).alignment = center_align
        ws.cell(row=3, column=start_col + daily_target_col_idx, value="Daily")
        for c in range(start_col, start_col + daily_target_col_idx + 1):
            cell = ws.cell(row=3, column=c)
            cell.font = table_header_font
            cell.fill = table_header_fill
            cell.alignment = center_align
            cell.border = thin_border

        ws.cell(row=4, column=start_col, value=li.end_date.strftime("%d %B, %Y"))
        ws.merge_cells(start_row=4, start_column=start_col + 1, end_row=4, end_column=start_col + 2)
        ws.cell(row=4, column=start_col + 1, value="Target impressions")
        ws.cell(row=4, column=start_col + daily_target_col_idx, value=target_impressions)

        ws.merge_cells(start_row=5, start_column=start_col + 1, end_row=5, end_column=start_col + 2)
        ws.cell(row=5, column=start_col + 1, value="Achieved impressions")
        impressions_col_letter = col_letters[1]
        ws.cell(
            row=5, column=start_col + daily_target_col_idx,
            value=f"=SUM({impressions_col_letter}{first_data_row}:{impressions_col_letter}{last_data_row})",
        )

        ws.merge_cells(start_row=6, start_column=start_col + 1, end_row=6, end_column=start_col + 2)
        ws.cell(row=6, column=start_col + 1, value="Remaining impressions")
        ws.cell(row=6, column=start_col + daily_target_col_idx, value=f"={col_d}4-{col_d}5")

        ws.merge_cells(start_row=7, start_column=start_col + 1, end_row=7, end_column=start_col + 2)
        ws.cell(row=7, column=start_col + 1, value="Daily Target")
        n_days = (last_data_row - first_data_row + 1) if last_data_row >= first_data_row else 0
        ws.cell(
            row=7, column=start_col + daily_target_col_idx,
            value=(
                f"=IFERROR(ROUND({col_d}6 / MAX(1, DATE({li.end_date.year},{li.end_date.month},{li.end_date.day}) - TODAY() + 1), 0), 0)"
                if n_days else 0
            ),
        )

        for r in range(4, 8):
            for c in range(start_col, start_col + daily_target_col_idx + 1):
                cell = ws.cell(row=r, column=c)
                cell.font = header_font if c == start_col + 1 else value_font
                if c == start_col + 1:
                    cell.fill = label_fill
                cell.border = thin_border
                cell.alignment = center_align

        ws.cell(row=8, column=start_col, value="Daily Report")
        ws.cell(
            row=8, column=start_col + 1,
            value=f"=SUM({impressions_col_letter}{first_data_row}:{impressions_col_letter}{last_data_row})",
        )
        clicks_col_letter = col_letters[2]
        ws.cell(
            row=8, column=start_col + 2,
            value=f"=SUM({clicks_col_letter}{first_data_row}:{clicks_col_letter}{last_data_row})",
        )
        ws.cell(
            row=8, column=start_col + 3,
            value=f"=IFERROR(({clicks_col_letter}8/{impressions_col_letter}8),0)",
        )
        for c in range(start_col, start_col + daily_target_col_idx + 1):
            cell = ws.cell(row=8, column=c)
            cell.font = header_font
            cell.fill = label_fill
            cell.border = thin_border
            cell.alignment = center_align

    # ── Non-video line items: side-by-side narrow blocks on the main sheet ──
    BLOCK_WIDTH = 6
    col_cursor = 1

    for li in non_video_line_items:
        start_col = col_cursor
        col_a = get_column_letter(start_col)
        col_b = get_column_letter(start_col + 1)
        col_c = get_column_letter(start_col + 2)
        col_d = get_column_letter(start_col + 3)

        ws_main.column_dimensions[col_a].width = 16
        ws_main.column_dimensions[col_b].width = 14
        ws_main.column_dimensions[col_c].width = 12
        ws_main.column_dimensions[col_d].width = 10
        for spacer_offset in range(4, BLOCK_WIDTH):
            ws_main.column_dimensions[get_column_letter(start_col + spacer_offset)].width = 10

        if not li.start_date or not li.end_date:
            col_cursor += BLOCK_WIDTH
            continue

        days = list(daterange(li.start_date, li.end_date))
        n_days = len(days)
        target_impressions = li.impressions or 0
        first_data_row = 10
        last_data_row = first_data_row + n_days - 1

        write_summary_header(
            ws_main, start_col, li, target_impressions, first_data_row, last_data_row,
            col_letters=[col_a, col_b, col_c, col_d], daily_target_col_idx=3,
        )

        ws_main.cell(row=9, column=start_col, value="Date")
        ws_main.cell(row=9, column=start_col + 1, value="Impressions")
        ws_main.cell(row=9, column=start_col + 2, value="Clicks")
        ws_main.cell(row=9, column=start_col + 3, value="CTR")
        for c in range(start_col, start_col + 4):
            cell = ws_main.cell(row=9, column=c)
            cell.font = header_font
            cell.fill = label_fill
            cell.alignment = center_align
            cell.border = thin_border

        li_entries = entries_map.get(li.id, {})
        for idx, day in enumerate(days):
            row = first_data_row + idx
            entry = li_entries.get(day)
            ws_main.cell(row=row, column=start_col, value=day.strftime("%d %B, %Y"))
            ws_main.cell(row=row, column=start_col + 1, value=entry.impressions if entry else 0)
            ws_main.cell(row=row, column=start_col + 2, value=entry.clicks if entry else 0)
            ws_main.cell(row=row, column=start_col + 3, value=entry.ctr if entry else 0)
            for c in range(start_col, start_col + 4):
                cell = ws_main.cell(row=row, column=c)
                cell.font = value_font
                cell.border = thin_border
                cell.alignment = center_align if c != start_col else left_align

        col_cursor += BLOCK_WIDTH

    # ── Video line items: one dedicated wide sheet each ──
    VIDEO_COLS = [
        "Date", "Impressions", "Clicks", "CTR", "Start Views", "Complete Views",
        "Video Completion Rate", "Viewable Impressions", "Measurable Impressions", "Viewability",
    ]
    VIDEO_COL_WIDTHS = [16, 14, 12, 10, 12, 14, 22, 18, 20, 14]

    for li in video_line_items:
        sheet_title = _safe_sheet_name(f"{li.line_item_id} - video", used_sheet_names)
        ws_v = wb.create_sheet(title=sheet_title)

        start_col = 1
        for offset, width in enumerate(VIDEO_COL_WIDTHS):
            ws_v.column_dimensions[get_column_letter(start_col + offset)].width = width

        if not li.start_date or not li.end_date:
            continue

        days = list(daterange(li.start_date, li.end_date))
        n_days = len(days)
        target_impressions = li.impressions or 0
        first_data_row = 10
        last_data_row = first_data_row + n_days - 1

        col_letters = [get_column_letter(start_col + i) for i in range(len(VIDEO_COLS))]
        write_summary_header(
            ws_v, start_col, li, target_impressions, first_data_row, last_data_row,
            col_letters=col_letters, daily_target_col_idx=3,
        )

        for i, col_name in enumerate(VIDEO_COLS):
            cell = ws_v.cell(row=9, column=start_col + i, value=col_name)
            cell.font = header_font
            cell.fill = label_fill
            cell.alignment = center_align
            cell.border = thin_border

        li_entries = entries_map.get(li.id, {})
        for idx, day in enumerate(days):
            row = first_data_row + idx
            entry = li_entries.get(day)
            row_values = [
                day.strftime("%d %B, %Y"),
                entry.impressions if entry else 0,
                entry.clicks if entry else 0,
                entry.ctr if entry else 0,
                entry.video_start if entry else 0,
                entry.video_end if entry else 0,
                entry.vcr if entry else 0,
                entry.viewable_impression if entry else 0,
                entry.measurable_impression if entry else 0,
                entry.viewability if entry else 0,
            ]
            for i, val in enumerate(row_values):
                cell = ws_v.cell(row=row, column=start_col + i, value=val)
                cell.font = value_font
                cell.border = thin_border
                cell.alignment = center_align if i != 0 else left_align

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()


# ── Add a single daily entry ────────────────────────────────────────────────

@api_view(["POST"])
def add_daily_entry(request, ticket_id):
    """
    Body: { "line_item_id": "LIBILL002", "date": "2026-05-15", "impressions": 50000, ... }
    Creates ONE row, never overwrites — 409 if an entry already exists for that date.
    """
    try:
        campaign = (
            Campaign.objects.select_related("insertion_order")
            .prefetch_related("line_items")
            .get(ticket_id=ticket_id)
        )
    except Campaign.DoesNotExist:
        return Response({"error": "Campaign not found"}, status=404)

    line_item_id = request.data.get("line_item_id")
    date_str = request.data.get("date")
    impressions = request.data.get("impressions")
    clicks = request.data.get("clicks")
    ctr = request.data.get("ctr")
    viewable_impression = request.data.get("viewable_impression")
    measurable_impression = request.data.get("measurable_impression")
    video_start = request.data.get("video_start")
    video_end = request.data.get("video_end")
    revenue = request.data.get("revenue")
    media_cost = request.data.get("media_cost")
    vcr = request.data.get("vcr")
    viewability = request.data.get("viewability")

    if not line_item_id or not date_str:
        return Response({"error": "line_item_id and date are required"}, status=400)

    try:
        line_item = LineItem.objects.get(line_item_id=line_item_id, campaign=campaign)
    except LineItem.DoesNotExist:
        return Response({"error": "LineItem not found in this campaign"}, status=404)

    entry_date = parse_date(date_str)
    if not entry_date:
        return Response({"error": "Invalid date"}, status=400)

    if entry_date < line_item.start_date or entry_date > line_item.end_date:
        return Response(
            {"error": f"Date must be between {line_item.start_date} and {line_item.end_date}"},
            status=400,
        )

    if CampaignDailyEntry.objects.filter(line_item=line_item, entry_date=entry_date).exists():
        return Response(
            {"error": "Entry already exists for this date. Editing is not allowed."}, status=409
        )

    def safe_num(v, cast=int):
        try:
            return cast(v)
        except (TypeError, ValueError):
            return 0

    entry = CampaignDailyEntry.objects.create(
        campaign=campaign,
        line_item=line_item,
        entry_date=entry_date,
        impressions=safe_num(impressions, int),
        clicks=safe_num(clicks, int),
        ctr=safe_num(ctr, float),
        viewable_impression=safe_num(viewable_impression, int),
        measurable_impression=safe_num(measurable_impression, int),
        video_start=safe_num(video_start, int),
        video_end=safe_num(video_end, int),
        revenue=safe_num(revenue, float),
        media_cost=safe_num(media_cost, float),
        vcr=safe_num(vcr, float),
        viewability=safe_num(viewability, float),
    )

    excel_bytes = build_daily_report_excel(campaign)
    # ✅ FIX: use ticket_id, not the undefined campaign_id
    filename = f"{ticket_id}_daily_report.xlsx"
    daily_excel_obj, _ = CampaignDailyReportExcel.objects.get_or_create(campaign=campaign)
    daily_excel_obj.excel_file.save(filename, ContentFile(excel_bytes), save=True)

    return Response(
        {
            "message": "Daily entry added successfully",
            "entry_id": entry.id,
            "line_item_id": line_item_id,
            "date": str(entry_date),
            "download_url": request.build_absolute_uri(daily_excel_obj.excel_file.url),
        },
        status=201,
    )


# ── Get already-submitted dates for a line item ─────────────────────────────

@api_view(["GET"])
def get_daily_entries(request, ticket_id):  # ✅ FIX: param name now matches urls.py
    line_item_id = request.query_params.get("line_item_id")

    try:
        campaign = Campaign.objects.get(ticket_id=ticket_id)  # ✅ FIX: ticket_id, not campaign_id
    except Campaign.DoesNotExist:
        return Response({"error": "Campaign not found"}, status=404)

    qs = CampaignDailyEntry.objects.filter(campaign=campaign)
    if line_item_id:
        qs = qs.filter(line_item__line_item_id=line_item_id)

    data = [
        {
            "line_item_id": e.line_item.line_item_id,
            "date": str(e.entry_date),
            "impressions": e.impressions,
            "clicks": e.clicks,
            "ctr": e.ctr,
            "viewable_impression": e.viewable_impression,
            "measurable_impression": e.measurable_impression,
            "video_start": e.video_start,
            "video_end": e.video_end,
            "revenue": e.revenue,
            "media_cost": e.media_cost,
            "vcr": e.vcr,
            "viewability": e.viewability,
        }
        for e in qs.order_by("entry_date")
    ]
    return Response(data)


# ── Generate/regenerate the daily-report Excel on demand ───────────────────

@api_view(["POST"])
def generate_daily_report_excel(request, ticket_id):  # ✅ FIX: param name
    try:
        campaign = (
            Campaign.objects.select_related("client", "insertion_order")
            .prefetch_related("line_items")
            .get(ticket_id=ticket_id)  # ✅ FIX
        )
    except Campaign.DoesNotExist:
        return Response({"error": "Campaign not found"}, status=404)

    excel_bytes = build_daily_report_excel(campaign)
    filename = f"{ticket_id}_daily_report.xlsx"

    daily_excel_obj, _ = CampaignDailyReportExcel.objects.get_or_create(campaign=campaign)
    daily_excel_obj.excel_file.save(filename, ContentFile(excel_bytes), save=True)

    return Response(
        {
            "message": "Daily report Excel generated successfully",
            "ticket_id": ticket_id,
            "download_url": request.build_absolute_uri(daily_excel_obj.excel_file.url),
            "filename": filename,
        },
        status=200,
    )


# ── Download the daily-report Excel ─────────────────────────────────────────

@api_view(["GET"])
def download_daily_report_excel(request, ticket_id):  # ✅ FIX: param name
    try:
        daily_excel_obj = CampaignDailyReportExcel.objects.get(campaign__ticket_id=ticket_id)  # ✅ FIX
    except CampaignDailyReportExcel.DoesNotExist:
        try:
            campaign = (
                Campaign.objects.select_related("client", "insertion_order")
                .prefetch_related("line_items")
                .get(ticket_id=ticket_id)  # ✅ FIX
            )
        except Campaign.DoesNotExist:
            return Response({"error": "Campaign not found"}, status=404)

        excel_bytes = build_daily_report_excel(campaign)
        response = HttpResponse(
            excel_bytes,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{ticket_id}_daily_report.xlsx"'
        return response

    response = FileResponse(
        daily_excel_obj.excel_file.open("rb"),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{ticket_id}_daily_report.xlsx"'
    return response


# ── Bulk upload from DV360-style Excel/CSV export ───────────────────────────

@api_view(["POST"])
@parser_classes([MultiPartParser])
def bulk_upload_daily_entries(request):
    """
    Upload an Excel/CSV file with columns (header names, case-insensitive):
      campaign (holds the ticket_id value), line item id (holds the dv_id value),
      date, impressions, clicks, click rate (ctr), revenue (adv currency),
      media cost (advertiser currency), start views, complete views,
      viewable impressions, measurable impressions
    """
    file = request.FILES.get("file")
    if not file:
        return Response({"error": "No file uploaded"}, status=400)

    try:
        wb = openpyxl.load_workbook(file, data_only=True)
        ws = wb.active
    except Exception:
        return Response({"error": "Invalid Excel file"}, status=400)

    HEADER_MAP = {
        "campaign": "ticket_id",  # ✅ this column holds the ticket_id value in this schema
        "line item id": "dv_id",
        "date": "date",
        "impressions": "impressions",
        "clicks": "clicks",
        "click rate (ctr)": "ctr",
        "revenue (adv currency)": "revenue",
        "media cost (advertiser currency)": "media_cost",
        "start views": "video_start",
        "complete views": "video_end",
        "viewable impressions": "viewable_impression",
        "measurable impressions": "measurable_impression",
    }

    raw_headers = [str(cell.value).strip().lower() if cell.value else "" for cell in ws[1]]

    col_idx = {}
    for i, h in enumerate(raw_headers):
        internal_key = HEADER_MAP.get(h)
        if internal_key:
            col_idx[internal_key] = i

    required_keys = ["ticket_id", "dv_id", "date", "impressions", "clicks", "ctr"]
    missing = [k for k in required_keys if k not in col_idx]
    if missing:
        reverse_map = {v: k for k, v in HEADER_MAP.items()}
        missing_labels = [reverse_map.get(k, k) for k in missing]
        return Response(
            {"error": f"Missing required column(s): {', '.join(missing_labels)}"}, status=400
        )

    def get_val(row, key):
        idx = col_idx.get(key)
        if idx is None:
            return None
        return row[idx].value

    results = {"inserted": [], "skipped": [], "errors": []}
    affected_campaigns = {}  # ticket_id → campaign object

    def safe_int(v):
        try:
            return int(float(str(v).replace(",", "").strip()))
        except Exception:
            return 0

    def safe_float(v):
        try:
            return float(str(v).replace(",", "").replace("%", "").strip())
        except Exception:
            return 0.0

    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=False), start=2):
        if all(cell.value is None for cell in row):
            continue

        ticket_id = str(get_val(row, "ticket_id") or "").strip()
        dv_id = str(get_val(row, "dv_id") or "").strip()
        date_raw = get_val(row, "date")
        impressions_raw = get_val(row, "impressions")
        clicks_raw = get_val(row, "clicks")
        ctr_raw = get_val(row, "ctr")
        revenue_raw = get_val(row, "revenue")
        media_cost_raw = get_val(row, "media_cost")
        video_start_raw = get_val(row, "video_start")
        video_end_raw = get_val(row, "video_end")
        viewable_imp_raw = get_val(row, "viewable_impression")
        measurable_imp_raw = get_val(row, "measurable_impression")

        if not ticket_id or not dv_id or not date_raw:
            results["errors"].append({
                "row": row_num,
                "reason": "Campaign, Line Item ID (DV ID), and Date are required",
                "data": f"{ticket_id} / {dv_id} / {date_raw}",
            })
            continue

        try:
            if hasattr(date_raw, "date"):
                entry_date = date_raw.date()
            else:
                entry_date = datetime.strptime(str(date_raw).strip(), "%Y-%m-%d").date()
        except Exception:
            results["errors"].append({
                "row": row_num,
                "reason": f"Invalid date format: '{date_raw}'. Use YYYY-MM-DD",
                "data": f"{ticket_id} / {dv_id}",
            })
            continue

        # ✅ FIX: no approval_status field on this Campaign model — ticket_id lookup only
        try:
            campaign = Campaign.objects.get(ticket_id=ticket_id)
        except Campaign.DoesNotExist:
            results["errors"].append({
                "row": row_num,
                "reason": f"Campaign '{ticket_id}' not found",
                "data": f"{ticket_id} / {dv_id}",
            })
            continue

        try:
            line_item = LineItem.objects.get(dv_id=dv_id, campaign=campaign)
        except LineItem.DoesNotExist:
            results["errors"].append({
                "row": row_num,
                "reason": f"Line item with DV ID '{dv_id}' not found in campaign '{ticket_id}'",
                "data": f"{ticket_id} / {dv_id}",
            })
            continue
        except LineItem.MultipleObjectsReturned:
            results["errors"].append({
                "row": row_num,
                "reason": f"Multiple line items share DV ID '{dv_id}' in campaign '{ticket_id}' — cannot determine which one",
                "data": f"{ticket_id} / {dv_id}",
            })
            continue

        if line_item.start_date and entry_date < line_item.start_date:
            results["errors"].append({
                "row": row_num,
                "reason": f"Date {entry_date} is before line item start date {line_item.start_date}",
                "data": f"{ticket_id} / {dv_id}",
            })
            continue

        if line_item.end_date and entry_date > line_item.end_date:
            results["errors"].append({
                "row": row_num,
                "reason": f"Date {entry_date} is after line item end date {line_item.end_date}",
                "data": f"{ticket_id} / {dv_id}",
            })
            continue

        if CampaignDailyEntry.objects.filter(line_item=line_item, entry_date=entry_date).exists():
            results["skipped"].append({
                "row": row_num,
                "reason": "Entry already exists for this date",
                "data": f"{ticket_id} / {dv_id} / {entry_date}",
            })
            continue

        CampaignDailyEntry.objects.create(
            campaign=campaign,
            line_item=line_item,
            entry_date=entry_date,
            impressions=safe_int(impressions_raw),
            clicks=safe_int(clicks_raw),
            ctr=safe_float(ctr_raw),
            viewable_impression=safe_int(viewable_imp_raw),
            measurable_impression=safe_int(measurable_imp_raw),
            video_start=safe_int(video_start_raw),
            video_end=safe_int(video_end_raw),
            revenue=safe_float(revenue_raw),
            media_cost=safe_float(media_cost_raw),
        )

        results["inserted"].append({"row": row_num, "data": f"{ticket_id} / {dv_id} / {entry_date}"})
        affected_campaigns[ticket_id] = campaign

    regenerated = []
    for tid, campaign in affected_campaigns.items():
        try:
            campaign_full = (
                Campaign.objects.select_related("insertion_order")
                .prefetch_related("line_items")
                .get(ticket_id=tid)
            )
            excel_bytes = build_daily_report_excel(campaign_full)
            filename = f"{tid}_daily_report.xlsx"
            daily_excel_obj, _ = CampaignDailyReportExcel.objects.get_or_create(campaign=campaign_full)
            daily_excel_obj.excel_file.save(filename, ContentFile(excel_bytes), save=True)
            regenerated.append(tid)
        except Exception:
            pass  # don't fail the whole response if one campaign's Excel regen fails

    return Response(
        {
            "message": "Bulk upload complete",
            "inserted": len(results["inserted"]),
            "skipped": len(results["skipped"]),
            "errors": len(results["errors"]),
            "details": results,
            "excel_regenerated_for": regenerated,
        },
        status=200,
    )
    
# ── Send Daily Report Excel via Email ───────────────────────────────────────

def daily_report_send_email_preview(request, ticket_id):
    campaign = get_object_or_404(Campaign, ticket_id=ticket_id)

    report_excel = getattr(campaign, "daily_report_excel", None)
    if not report_excel or not report_excel.excel_file:
        return JsonResponse({
            "status": False,
            "message": "No daily report Excel generated yet for this campaign.",
        })

    client = campaign.client
    to_emails = [e.strip() for e in (client.default_email_send_to or "").split(",") if e.strip()]
    cc_emails = [e.strip() for e in (client.default_email_send_cc or "").split(",") if e.strip()]

    if not to_emails:
        return JsonResponse({
            "status": False,
            "message": f"No 'Default Email To' configured for {client.name}.",
            "needs_email_setup": True,
            "company_id": client.id,
        })

    subject = f"Daily Report - {campaign.campaign_name} ({ticket_id})"
    filename = report_excel.excel_file.name.split("/")[-1]

    return JsonResponse({
        "status": True,
        "ticket_id": ticket_id,
        "campaign_name": campaign.campaign_name,
        "client_name": client.name,
        "subject": subject,
        "to": to_emails,
        "cc": cc_emails,
        "attachments": [filename],
    })


@csrf_exempt
def daily_report_send_email(request, ticket_id):
    if request.method != "POST":
        return JsonResponse({"status": False, "message": "Invalid request method."})

    campaign = get_object_or_404(Campaign, ticket_id=ticket_id)

    report_excel = getattr(campaign, "daily_report_excel", None)
    if not report_excel or not report_excel.excel_file:
        return JsonResponse({"status": False, "message": "No daily report Excel generated yet."})

    client = campaign.client
    to_emails = [e.strip() for e in (client.default_email_send_to or "").split(",") if e.strip()]
    cc_emails = [e.strip() for e in (client.default_email_send_cc or "").split(",") if e.strip()]

    if not to_emails:
        return JsonResponse({"status": False, "message": "No 'Default Email To' configured."})

    subject = request.POST.get("subject") or f"Daily Report - {campaign.campaign_name}"
    body_html = request.POST.get("email_body") or (
        f"Please find attached the daily report for {campaign.campaign_name}."
    )

    email = EmailMultiAlternatives(
        subject=subject, body="", from_email=settings.DEFAULT_FROM_EMAIL,
        to=to_emails, cc=cc_emails if cc_emails else None,
    )
    email.attach_alternative(body_html, "text/html")

    try:
        filename = report_excel.excel_file.name.split("/")[-1]
        report_excel.excel_file.open("rb")
        email.attach(
            filename,
            report_excel.excel_file.read(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        email.send(fail_silently=False)
    except Exception as e:
        return JsonResponse({"status": False, "message": f"Failed to send email: {e}"})
    finally:
        report_excel.excel_file.close()

    return JsonResponse({"status": True, "message": f"Email sent to {', '.join(to_emails)}"})